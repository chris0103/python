from docx import Document

doc = Document('周4组.docx')
for no, doc_table in enumerate(doc.tables):
    for i in range(0, len(doc_table.rows)):
        print(i, doc_table.cell(i, 1).text)


