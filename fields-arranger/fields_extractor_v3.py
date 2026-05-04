import openpyxl
import shutil
import os

from openpyxl.worksheet.worksheet import Worksheet


class Field:
    """地块"""

    def __init__(self, no, name, east_to, south_to, west_to, north_to, contract_area, actual_area):
        self.no = no
        self.name = name
        self.east_to = east_to
        self.south_to = south_to
        self.west_to = west_to
        self.north_to = north_to
        self.contract_area = contract_area
        self.actual_area = actual_area

    def __repr__(self):
        return f'{self.no}\t{self.name}\t{self.east_to}\t{self.south_to}\t{self.west_to}\t{self.north_to}\t{self.contract_area}\t{self.actual_area}\t'


class Person:
    """人员"""

    def __init__(self, name):
        self.name = name
        self.fields = []

    def add_field(self, field):
        self.fields.append(field)

    @property
    def total_contract_area(self):
        return sum(field.contract_area for field in self.fields)

    @property
    def total_actual_area(self):
        return sum(field.actual_area for field in self.fields)

    def __repr__(self):
        return f'{self.name}\t{self.fields}\nTotal contract area: {self.total_contract_area}\tTotal actual area: {self.total_actual_area}\n'


def rectify_area(area: str) -> str:
    if area is not None and area != '':
        area = area.replace('。', '.').replace('\n', '').replace(',', '')
    return area


def process_sheet(sheet: Worksheet, start_column: int = 2):
    print(sheet_name)
    current_name = ''
    for row in range(5, sheet.max_row):
        column = start_column
        name = sheet.cell(row, column).value
        column += 3
        field_name = sheet.cell(row, column).value
        column += 1
        field_no = sheet.cell(row, column).value
        column += 1
        east_to = sheet.cell(row, column).value
        column += 1
        south_to = sheet.cell(row, column).value
        column += 1
        west_to = sheet.cell(row, column).value
        column += 1
        north_to = sheet.cell(row, column).value
        column += 1
        contract_area = sheet.cell(row, column).value
        # print(type(contract_area))
        if type(contract_area) is str:
            contract_area = rectify_area(contract_area)
        column += 1
        actual_area = sheet.cell(row, column).value
        # rint(type(actual_area))
        if type(actual_area) is str:
            actual_area = rectify_area(actual_area)

        if name is not None and name != '' and name != current_name:
            """new person"""
            current_name = name
            person = Person(current_name)
            persons.append(person)

        if field_no is not None and field_no != '':  # ignore empty rows
            field = Field(field_no, field_name, east_to, south_to, west_to, north_to, float(contract_area),
                          float(actual_area))
            person.add_field(field)


def generate_table(village: str, group: str):
    filename = f'{village}{group}.xlsx'
    shutil.copy("files/template.xlsx", f'{filename}')
    wb_table = openpyxl.load_workbook(f'{filename}')
    template_sheet = wb_table.worksheets[0]
    for person in persons:
        ws = wb_table.copy_worksheet(template_sheet)
        ws.title = person.name
        ws.cell(row=3, column=1, value=f'兴东镇 {village}村 {group}')
        ws.cell(row=4, column=3, value=person.name)
        ws.cell(row=4, column=8, value=phones.get(person.name))
        ws.cell(row=5, column=3, value=len(person.fields))
        ws.cell(row=5, column=8, value=person.total_contract_area)
        ws.cell(row=5, column=11, value=person.total_actual_area)

        for idx, field in enumerate(person.fields):
            row = 6 + idx * 3
            ws.cell(row=row, column=3, value=field.no)
            ws.cell(row=row, column=12, value=field.actual_area)
            row += 1
            east_to = ws.cell(row=row, column=4, value=field.east_to)
            ws.cell(row=row, column=6, value=field.south_to)
            ws.cell(row=row + 1, column=4, value=field.west_to)
            ws.cell(row=row + 1, column=6, value=field.north_to)

    del wb_table['Sheet1']
    wb_table.save(f'{filename}')


# 获取通讯录
phones = {}

# 格式一
# wb_phones = openpyxl.load_workbook('files/周家通讯录.xlsx')
# sheet = wb_phones.worksheets[0]
# for row in range(4, sheet.max_row):
#     name = sheet.cell(row, 2).value
#     phone = sheet.cell(row, 3).value
#     phones[name] = phone

# 格式二
wb_phones = openpyxl.load_workbook('files/周家通讯录.xlsx')
for i in range(0, len(wb_phones.sheetnames) - 1):
    sheet = wb_phones.worksheets[i]
    for row in range(3, sheet.max_row):
        name = sheet.cell(row, 3).value
        phone = sheet.cell(row, 11).value
        phones[name] = phone

# 抽取信息

# 格式一
# directory = os.fsencode("files/牛陆")
# for file in os.listdir(directory):
#     filename = os.fsdecode(file)
#     if not filename.endswith(".xlsx"):
#         continue
#     village = f'{filename[0:1]}家'
#     group = filename[1:len(filename) - 5]
#     print(f'{village}{group}')
#     wb_fields = openpyxl.load_workbook(f'files/牛陆/{filename}')
#     persons = []
#     sheet = wb_fields.worksheets[0]
#     process_sheet(sheet)
#     generate_table()


# 格式二
village = '周家庄'
wb_fields = openpyxl.load_workbook(f'files/{village}.xlsx')
for i in range(0, len(wb_fields.sheetnames)):
    persons = []
    sheet_name = wb_fields.sheetnames[i]
    sheet = wb_fields.worksheets[i]
    process_sheet(sheet, 2)
    generate_table(village, sheet_name)
