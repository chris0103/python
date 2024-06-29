import openpyxl
from openpyxl import Workbook


class Field:
    """地块"""

    def __init__(self, no, name, east_to, south_to, west_to, north_to, contract_area, actual_area):
        self.no = no
        self.name = name
        self.east_to = east_to
        self.south_to = south_to
        self.west_to = west_to
        self.north_to = north_to
        self.contract_area = contract_area
        self.actual_area = actual_area

    def __repr__(self):
        return f'{self.no}\t{self.name}\t{self.east_to}\t{self.south_to}\t{self.west_to}\t{self.north_to}\t{self.contract_area}\t{self.actual_area}\t'


class Person:
    """人员"""

    def __init__(self, name):
        self.name = name
        self.fields = []

    def add_field(self, field):
        self.fields.append(field)

    @property
    def total_contract_area(self):
        return sum(field.contract_area for field in self.fields)

    @property
    def total_actual_area(self):
        return sum(field.actual_area for field in self.fields)

    def __repr__(self):
        return f'{self.name}\t{self.fields}\nTotal contract area: {self.total_contract_area}\tTotal actual area: {self.total_actual_area}\n'


phones = {}
wb_phones = openpyxl.load_workbook('phones.xlsx')
sheet = wb_phones.worksheets[0]
for row in range(4, sheet.max_row):
    name = sheet.cell(row, 2).value
    phone = sheet.cell(row, 3).value
    phones[name] = phone
# print(phones)


wb_fields = openpyxl.load_workbook('fields.xlsx')

# get sheet names
# print(wb.sheetnames)

persons = []
person = None
current_name = ''
sheet = wb_fields.worksheets[0]
for row in range(5, sheet.max_row):
    name = sheet.cell(row, 2).value
    field_name = sheet.cell(row, 5).value
    field_no = sheet.cell(row, 6).value
    east_to = sheet.cell(row, 7).value
    south_to = sheet.cell(row, 8).value
    west_to = sheet.cell(row, 9).value
    north_to = sheet.cell(row, 10).value
    contract_area = sheet.cell(row, 11).value
    actual_area = sheet.cell(row, 12).value

    if name is not None:
        """new person"""
        current_name = name
        person = Person(current_name)
        persons.append(person)

    if field_no is not None:  # ignore empty rows
        field = Field(field_no, field_name, east_to, south_to, west_to, north_to, contract_area, actual_area)
        person.add_field(field)

# print persons
# for person in persons:
#    print(person)

wb_table = Workbook()
for person in persons:
    ws = wb_table.create_sheet(person.name)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    ws.cell(row=1, column=1, value=f'附件3')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)
    ws.cell(row=2, column=1, value=f'兴化市“小田变大田”改革农户基础情况调查表')
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=12)
    ws.cell(row=3, column=1, value=f'xx镇 xx村 xx组')

    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=2)
    ws.cell(row=4, column=1, value=f'承包方姓名')
    ws.merge_cells(start_row=4, start_column=3, end_row=4, end_column=6)
    ws.cell(row=4, column=3, value=person.name)
    ws.cell(row=4, column=7, value=f'联系方式')
    ws.merge_cells(start_row=4, start_column=8, end_row=4, end_column=12)
    ws.cell(row=4, column=8, value=phones.get(person.name))

    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=2)
    ws.cell(row=5, column=1, value=f'确权地块数')
    ws.merge_cells(start_row=5, start_column=3, end_row=5, end_column=6)
    ws.cell(row=5, column=3, value=len(person.fields))
    ws.cell(row=5, column=7, value=f'承包面积')
    ws.merge_cells(start_row=5, start_column=8, end_row=5, end_column=9)
    ws.cell(row=5, column=8, value=person.total_contract_area)
    ws.cell(row=5, column=10, value=f'实测面积')
    ws.merge_cells(start_row=5, start_column=11, end_row=5, end_column=12)
    ws.cell(row=5, column=11, value=person.total_actual_area)

    for idx, field in enumerate(person.fields):
        row = 6 + idx * 3
        ws.merge_cells(start_row=row, start_column=1, end_row=row+2, end_column=1)
        ws.cell(row=row, column=1, value=f'地块{idx + 1}')
        ws.cell(row=row, column=2, value=f'地块编码')
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        ws.cell(row=row, column=3, value=field.no)
        ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=9)
        ws.cell(row=row, column=7, value=f'是否愿意流转土地经营权')
        ws.cell(row=row, column=11, value='实测面积')
        ws.cell(row=row, column=12, value=field.actual_area)
        row += 1
        ws.merge_cells(start_row=row, start_column=2, end_row=row+1, end_column=2)
        ws.cell(row=row, column=2, value=f'四至')
        ws.cell(row=row, column=3, value=f'东至')
        ws.cell(row=row, column=4, value=field.east_to)
        ws.cell(row=row, column=5, value=f'南至')
        ws.cell(row=row, column=6, value=field.south_to)
        ws.cell(row=row+1, column=3, value=f'西至')
        ws.cell(row=row+1, column=4, value=field.west_to)
        ws.cell(row=row+1, column=5, value=f'北至')
        ws.cell(row=row+1, column=6, value=field.north_to)
        ws.merge_cells(start_row=row, start_column=7, end_row=row+1, end_column=7)
        ws.cell(row=row, column=7, value=f'平台流转面积')
        ws.merge_cells(start_row=row, start_column=8, end_row=row+1, end_column=8)
        ws.merge_cells(start_row=row, start_column=9, end_row=row + 1, end_column=9)
        ws.cell(row=row, column=9, value=f'私下流转面积')
        ws.merge_cells(start_row=row, start_column=10, end_row=row + 1, end_column=10)
        ws.merge_cells(start_row=row, start_column=11, end_row=row + 1, end_column=11)
        ws.cell(row=row, column=11, value=f'征用面积')

    row += 3
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row=row, column=1, value='是否愿意参加“小田变大田”改革')
    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=12)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row=row, column=1, value='对继续种植经营的地块有何意见或建议')
    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=12)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row=row, column=1, value='对流转土地经营权有何意见或建议')
    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=12)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row=row, column=1, value='其它需要反映的建议')
    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=12)

del wb_table['Sheet']
wb_table.save('stats.xlsx')





# get cells area
# print(sheet.dimensions)

# get row count and column count
# print(sheet.max_row, sheet.max_column)

# get value in specified cell
# print(sheet.cell(5, 2).value)
# print(sheet['B5'].value)

# get a matrix of cells (in a 2-dimension array)
# print(sheet['G11:I14'])
