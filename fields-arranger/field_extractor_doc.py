from docx import Document
import openpyxl
import shutil
import os

class Field:
    """地块"""

    def __init__(self, no, name, east_to, south_to, west_to, north_to, contract_area, actual_area):
        self.no = no
        self.name = name
        self.east_to = east_to
        self.south_to = south_to
        self.west_to = west_to
        self.north_to = north_to
        self.contract_area = contract_area
        self.actual_area = actual_area

    def __repr__(self):
        return f'{self.no}\t{self.name}\t{self.east_to}\t{self.south_to}\t{self.west_to}\t{self.north_to}\t{self.contract_area}\t{self.actual_area}\t'


class Person:
    """人员"""

    def __init__(self, name):
        self.name = name
        self.fields = []

    def add_field(self, field):
        self.fields.append(field)

    @property
    def total_contract_area(self):
        return sum(field.contract_area for field in self.fields)

    @property
    def total_actual_area(self):
        return sum(field.actual_area for field in self.fields)

    def __repr__(self):
        return f'{self.name}\t{self.fields}\nTotal contract area: {self.total_contract_area}\tTotal actual area: {self.total_actual_area}\n'

phones = {}

# 格式一
# wb_phones = openpyxl.load_workbook('files/西鲍通讯录.xlsx')
# sheet = wb_phones.worksheets[0]
# for row in range(4, sheet.max_row):
#     name = sheet.cell(row, 2).value
#     phone = sheet.cell(row, 3).value
#     phones[name] = phone

# 格式二
wb_phones = openpyxl.load_workbook('files/周家通讯录.xlsx')
for i in range(0, len(wb_phones.sheetnames) - 1):
    sheet = wb_phones.worksheets[i]
    for row in range(3, sheet.max_row):
        name = sheet.cell(row, 3).value
        phone = sheet.cell(row, 11).value
        phones[name] = phone

directory = os.fsencode("files/周家docx")
for file in os.listdir(directory):
    filename = os.fsdecode(file)
    if not filename.endswith(".docx"):
        continue
    village= f'{filename[0:1]}家'
    group = filename[1:len(filename) - 5]
    print(f'{village}{group}')

    doc = Document(f'files/周家docx/{filename}')

    persons = []
    person = None
    current_name = ''

    for no, doc_table in enumerate(doc.tables):
        row_count = len(doc_table.rows)
        if row_count < 2:
            continue

        for row in range(2, row_count):
            name = doc_table.cell(row, 1).text
            field_name = doc_table.cell(row, 4).text
            field_no = doc_table.cell(row, 5).text
            east_to = doc_table.cell(row, 6).text
            south_to = doc_table.cell(row, 7).text
            west_to = doc_table.cell(row, 8).text
            north_to = doc_table.cell(row, 9).text
            contract_area = doc_table.cell(row, 10).text.replace('。', '.').replace('\n', '')
            actual_area = doc_table.cell(row, 11).text.replace('。', '.').replace('\n', '')

            if name is not None and name != '' and name != current_name:
                """new person"""
                current_name = name
                person = Person(current_name)
                persons.append(person)

            if field_no is not None and field_no != '':  # ignore empty rows
                field = Field(field_no, field_name, east_to, south_to, west_to, north_to, float(contract_area), float(actual_area))
                person.add_field(field)

    filename = filename.replace("docx", "xlsx")
    shutil.copy("files/template.xlsx", f'{filename}')
    wb_table = openpyxl.load_workbook(f'{filename}')
    template_sheet = wb_table.worksheets[0]
    for person in persons:
        ws = wb_table.copy_worksheet(template_sheet)
        ws.title = person.name
        ws.cell(row=3, column=1, value=f'兴东镇 {village}村 {group}')
        ws.cell(row=4, column=3, value=person.name)
        ws.cell(row=4, column=8, value=phones.get(person.name))
        ws.cell(row=5, column=3, value=len(person.fields))
        ws.cell(row=5, column=8, value=person.total_contract_area)
        ws.cell(row=5, column=11, value=person.total_actual_area)

        for idx, field in enumerate(person.fields):
            row = 6 + idx * 3
            ws.cell(row=row, column=3, value=field.no)
            ws.cell(row=row, column=12, value=field.actual_area)
            row += 1
            east_to = ws.cell(row=row, column=4, value=field.east_to)
            ws.cell(row=row, column=6, value=field.south_to)
            ws.cell(row=row + 1, column=4, value=field.west_to)
            ws.cell(row=row + 1, column=6, value=field.north_to)


    del wb_table['Sheet1']
    wb_table.save(f'{filename}')
